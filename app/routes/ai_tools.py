"""
AI Tools Routes - API endpoints for AI-powered tools
Available to all authenticated users with role-based tool access
"""
from flask import Blueprint, jsonify, request, session
from functools import wraps

ai_tools_bp = Blueprint('ai_tools', __name__, url_prefix='/api/ai-tools')


def is_app_frozen():
    """Check if the application is in frozen/maintenance mode"""
    try:
        from app.models import SiteSettings
        settings = SiteSettings.query.first()
        return settings.app_frozen if settings else False
    except Exception:
        return False


def get_freeze_message(lang='en'):
    """Get the freeze message in the specified language"""
    try:
        from app.models import SiteSettings
        settings = SiteSettings.query.first()
        if not settings:
            return 'The platform is currently under maintenance.'
        return settings.freeze_message_ar if lang == 'ar' else settings.freeze_message
    except Exception:
        return 'The platform is currently under maintenance.'


def login_required(f):
    """Decorator to require user authentication"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return jsonify({'error': 'Authentication required'}), 401
        return f(*args, **kwargs)
    return decorated_function


@ai_tools_bp.route('/list', methods=['GET'])
@login_required
def get_tools_list():
    """Get list of available AI tools for current user's role"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        role = session.get('role', 'student')
        tools = ai_tools_service.get_tools_for_role(role)
        
        return jsonify({
            'success': True,
            'tools': tools,
            'role': role
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/prompts/<tool_id>', methods=['GET'])
@login_required
def get_prompt_library(tool_id):
    """Return the ready-made starter prompts for a given tool."""
    try:
        from app.services.ai_tools_service import (
            ai_tools_service,
            get_prompt_library_for,
        )

        role = session.get('role', 'student')
        available_tools = ai_tools_service.get_tools_for_role(role)
        if tool_id not in [t['id'] for t in available_tools]:
            return jsonify({'error': f'Tool "{tool_id}" is not available for your role'}), 403

        return jsonify({
            'success': True,
            'tool_id': tool_id,
            'prompts': get_prompt_library_for(tool_id)
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/all', methods=['GET'])
@login_required
def get_all_tools():
    """Get all tools organized by category (admin only)"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        role = session.get('role', '').lower()
        if role not in ['admin', 'super_admin', 'superadmin', 'institution_admin']:
            return jsonify({'error': 'Access denied'}), 403
        
        all_tools = ai_tools_service.get_all_tools()
        
        return jsonify({
            'success': True,
            'tools': all_tools
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/execute', methods=['POST'])
@login_required
def execute_tool():
    """Execute an AI tool with given parameters"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json()
        if not data:
            return jsonify({'error': 'Request body required'}), 400
        
        tool_id = data.get('tool_id')
        if not tool_id:
            return jsonify({'error': 'tool_id is required'}), 400
        
        params = data.get('params', {})
        model = data.get('model', 'gpt-4.1')
        
        role = session.get('role', 'student')
        available_tools = ai_tools_service.get_tools_for_role(role)
        tool_ids = [t['id'] for t in available_tools]
        
        if tool_id not in tool_ids:
            return jsonify({'error': f'Tool "{tool_id}" is not available for your role'}), 403
        
        result = ai_tools_service.execute_tool(tool_id, params, model)
        
        if result.get('success'):
            return jsonify(result)
        else:
            return jsonify(result), 500
            
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/text-generator', methods=['POST'])
@login_required
def text_generator():
    """Generate text based on user input"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json() or {}
        params = {
            'input': data.get('topic', data.get('input', '')),
            'tone': data.get('tone', 'professional'),
            'length': data.get('length', 'medium'),
            'language': data.get('language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('text_generator', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/summarizer', methods=['POST'])
@login_required
def summarizer():
    """Summarize text"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json() or {}
        params = {
            'input': data.get('text', data.get('input', '')),
            'summary_type': data.get('summary_type', 'key points'),
            'length': data.get('length', 'medium'),
            'language': data.get('language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('summarizer', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/translator', methods=['POST'])
@login_required
def translator():
    """Translate text"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json() or {}
        params = {
            'input': data.get('text', data.get('input', '')),
            'source_lang': data.get('source_language', 'auto-detect'),
            'target_lang': data.get('target_language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('translator', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/grammar-checker', methods=['POST'])
@login_required
def grammar_checker():
    """Check grammar and style"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json() or {}
        params = {
            'input': data.get('text', data.get('input', '')),
            'language': data.get('language', 'English'),
            'style': data.get('style', 'formal')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('grammar_checker', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/paraphraser', methods=['POST'])
@login_required
def paraphraser():
    """Paraphrase text"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json() or {}
        params = {
            'input': data.get('text', data.get('input', '')),
            'tone': data.get('tone', 'professional'),
            'style': data.get('style', 'formal'),
            'language': data.get('language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('paraphraser', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/lesson-planner', methods=['POST'])
@login_required
def lesson_planner():
    """Generate lesson plan"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        role = session.get('role', '').lower()
        if role not in ['teacher', 'instructor', 'admin', 'super_admin', 'superadmin', 'institution_admin']:
            return jsonify({'error': 'This tool is for teachers only'}), 403
        
        data = request.get_json() or {}
        params = {
            'topic': data.get('topic', ''),
            'subject': data.get('subject', 'general'),
            'level': data.get('level', 'intermediate'),
            'duration': data.get('duration', '60 minutes'),
            'language': data.get('language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('lesson_planner', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/question-generator', methods=['POST'])
@login_required
def question_generator():
    """Generate exam/quiz questions"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        role = session.get('role', '').lower()
        if role not in ['teacher', 'instructor', 'admin', 'super_admin', 'superadmin', 'institution_admin']:
            return jsonify({'error': 'This tool is for teachers only'}), 403
        
        data = request.get_json() or {}
        params = {
            'topic': data.get('topic', ''),
            'subject': data.get('subject', 'general'),
            'difficulty': data.get('difficulty', 'mixed'),
            'question_types': data.get('question_types', 'multiple choice, true/false'),
            'num_questions': str(data.get('num_questions', 5)),
            'language': data.get('language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('question_generator', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/study-assistant', methods=['POST'])
@login_required
def study_assistant():
    """Get study help"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json() or {}
        params = {
            'input': data.get('question', data.get('input', '')),
            'subject': data.get('subject', 'general'),
            'level': data.get('level', 'intermediate'),
            'language': data.get('language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('study_assistant', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/flashcard-generator', methods=['POST'])
@login_required
def flashcard_generator():
    """Generate study flashcards"""
    try:
        from app.services.ai_tools_service import ai_tools_service
        
        data = request.get_json() or {}
        params = {
            'topic': data.get('topic', ''),
            'input': data.get('content', data.get('input', '')),
            'num_cards': str(data.get('num_cards', 10)),
            'language': data.get('language', 'English')
        }
        model = data.get('model', 'gpt-4.1')
        
        result = ai_tools_service.execute_tool('flashcard_generator', params, model)
        return jsonify(result)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/my-documents', methods=['GET'])
@login_required
def get_my_documents():
    """Get documents from courses the student is enrolled in"""
    import os
    import json
    
    user_id = session.get('user_id')
    role = session.get('role', '').lower()
    
    try:
        # Load course files
        course_files_path = 'course_files.json'
        if not os.path.exists(course_files_path):
            return jsonify({'success': True, 'documents': []})
        
        with open(course_files_path, 'r') as f:
            course_files_data = json.load(f)
        
        # Load classes to get course info
        classes_path = 'classes.json'
        classes_data = {}
        if os.path.exists(classes_path):
            with open(classes_path, 'r') as f:
                classes_data = json.load(f)
        
        # Get user's enrolled classes
        enrolled_class_ids = set()
        for class_id, class_info in classes_data.items():
            students = class_info.get('students', [])
            if user_id in students or role in ['admin', 'super_admin', 'superadmin', 'teacher', 'instructor']:
                enrolled_class_ids.add(class_id)
        
        # Filter documents to only those from enrolled courses
        documents = []
        for file_info in course_files_data.get('files', []):
            if file_info.get('class_id') in enrolled_class_ids:
                # Only include readable document types
                file_type = file_info.get('file_type', '').lower()
                if file_type in ['pdf', 'doc', 'docx', 'txt', 'ppt', 'pptx']:
                    class_name = classes_data.get(file_info.get('class_id'), {}).get('name', 'Unknown Course')
                    documents.append({
                        'file_id': file_info.get('file_id'),
                        'filename': file_info.get('original_filename'),
                        'file_type': file_type,
                        'class_id': file_info.get('class_id'),
                        'class_name': class_name,
                        'description': file_info.get('description', ''),
                        'uploaded_at': file_info.get('uploaded_at', '')
                    })
        
        return jsonify({'success': True, 'documents': documents})
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@ai_tools_bp.route('/content-chat', methods=['POST'])
@login_required
def content_chat():
    """Chat with document content"""
    import os
    import json
    
    if is_app_frozen():
        role = session.get('role', '').lower()
        if role not in ['superadmin', 'super_admin', 'admin']:
            data = request.get_json() or {}
            lang = data.get('language', 'en')[:2]
            return jsonify({'error': get_freeze_message(lang), 'app_frozen': True}), 503
    
    try:
        from app.services.ai_tools_service import ai_tools_service
        from app.utils.file_handler import FileHandler
        
        data = request.get_json() or {}
        file_id = data.get('file_id')
        question = data.get('question', data.get('input', ''))
        language = data.get('language', 'English')
        model = data.get('model', 'gpt-4.1')
        
        if not file_id:
            return jsonify({'error': 'Please select a document'}), 400
        
        if not question:
            return jsonify({'error': 'Please enter a question'}), 400
        
        # Load file metadata
        course_files_path = 'course_files.json'
        if not os.path.exists(course_files_path):
            return jsonify({'error': 'No documents available'}), 404
        
        with open(course_files_path, 'r') as f:
            course_files_data = json.load(f)
        
        # Find the file
        file_info = None
        for f_info in course_files_data.get('files', []):
            if f_info.get('file_id') == file_id:
                file_info = f_info
                break
        
        if not file_info:
            return jsonify({'error': 'Document not found'}), 404
        
        # Extract text from file
        file_path = os.path.join('course_files', file_info.get('stored_filename', ''))
        if not os.path.exists(file_path):
            return jsonify({'error': 'Document file not found on server'}), 404
        
        file_type = file_info.get('file_type', '').lower()
        document_content = ""
        
        if file_type in ['pdf', 'doc', 'docx', 'txt']:
            document_content = FileHandler.extract_text(file_path)
        elif file_type in ['ppt', 'pptx']:
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                text_runs = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text_runs.append(shape.text)
                document_content = "\n".join(text_runs)[:15000]
            except Exception as e:
                document_content = f"Could not extract text from presentation: {str(e)}"
        
        if not document_content or len(document_content) < 50:
            return jsonify({'error': 'Could not extract readable content from this document'}), 400
        
        # Prepare params for the AI tool
        params = {
            'input': question,
            'document_title': file_info.get('original_filename', 'Unknown Document'),
            'document_content': document_content[:12000],  # Limit content size
            'language': language
        }
        
        result = ai_tools_service.execute_tool('content_chat', params, model)
        return jsonify(result)
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500
