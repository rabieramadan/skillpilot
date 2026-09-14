"""
PresentMate - Multi-AI Collaborative Presentation Generator
Generates slides from text or files and lets users convert to different formats
"""

import json
import os
import base64
from typing import Dict, Any, List
from datetime import datetime

# Try to import python-pptx (optional dependency)
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False
    Presentation = None


class PresentationService:
    """
    Service for generating presentations from text/files
    """
    
    def __init__(self, ai_service):
        self.ai_service = ai_service
        self.sessions = {}
        
    def generate_slides_content(self, message: str, files: List[str] = None) -> Dict[str, Any]:
        """
        Generate slides content from user's text or files
        Returns the slides content to show in the prompt
        """
        session_id = str(hash(message + str(datetime.now().timestamp())))
        
        try:
            # Get API key from Config (same way as other services)
            from config.config import Config
            config = Config()
            api_key = config.OPENAI_API_KEY
            
            if not api_key or api_key == 'sk-your-openai-key-here':
                return {
                    'status': 'error',
                    'error': 'OpenAI API key not configured',
                    'message': 'Please configure your OpenAI API key to use PresentMate'
                }
            
            # Create prompt for slide generation
            prompt = f"""Based on the following input, create a presentation outline with slides.

User Input: {message}

Generate a well-structured presentation with:
1. A clear title
2. Between 5-10 slides
3. Each slide should have:
   - A descriptive title
   - 2-4 key bullet points
   - A brief visual suggestion

Format your response as JSON:
{{
  "title": "Presentation Title",
  "slides": [
    {{
      "number": 1,
      "title": "Slide Title",
      "content": ["Point 1", "Point 2", "Point 3"],
      "visual_suggestion": "Description of suggested visual"
    }}
  ]
}}

Make it professional, clear, and engaging."""

            # Call AI service to generate content
            response = self.ai_service.chat(
                provider='openai',
                message=prompt,
                api_key=api_key,
                files=files,
                conversation_history=[],
                version='gpt-4-turbo'
            )
            
            # Parse the response
            text = response.get('text', '{}')
            
            # Extract JSON from response
            start_idx = text.find('{')
            end_idx = text.rfind('}') + 1
            if start_idx != -1 and end_idx > start_idx:
                json_text = text[start_idx:end_idx]
                slides_data = json.loads(json_text)
            else:
                return {
                    'status': 'error',
                    'error': 'Failed to parse AI response',
                    'message': 'Could not generate slides content'
                }
            
            # Store the session data
            self.sessions[session_id] = {
                'slides_data': slides_data,
                'created_at': datetime.now().isoformat()
            }
            
            # Format the slides for display
            slides_text = self._format_slides_for_display(slides_data)
            
            return {
                'status': 'generated',
                'session_id': session_id,
                'slides_text': slides_text,
                'slides_data': slides_data,
                'message': 'Slides content generated! You can now convert to your preferred format.',
                'conversion_prompt': 'Would you like to convert this to PowerPoint (.pptx), HTML, or download as JSON?'
            }
            
        except Exception as e:
            return {
                'status': 'error',
                'error': str(e),
                'message': f'Failed to generate slides: {str(e)}'
            }
    
    def _format_slides_for_display(self, slides_data: Dict[str, Any]) -> str:
        """Format slides data as readable text"""
        title = slides_data.get('title', 'Untitled Presentation')
        slides = slides_data.get('slides', [])
        
        output = [f"📊 **{title}**\n"]
        output.append(f"Total Slides: {len(slides)}\n")
        output.append("="*50 + "\n\n")
        
        for slide in slides:
            slide_num = slide.get('number', 0)
            slide_title = slide.get('title', 'Untitled')
            content = slide.get('content', [])
            visual = slide.get('visual_suggestion', 'No visual')
            
            output.append(f"**Slide {slide_num}: {slide_title}**\n")
            for point in content:
                output.append(f"  • {point}\n")
            output.append(f"  🎨 Visual: {visual}\n\n")
        
        return "".join(output)
    
    def _generate_images_for_slides(self, slides_data: Dict[str, Any], session_id: str) -> Dict[str, Any]:
        """Generate DALL-E images for each slide based on visual suggestions"""
        try:
            from config.config import Config
            config = Config()
            api_key = config.OPENAI_API_KEY
            
            if not api_key or api_key == 'sk-your-openai-key-here':
                return {
                    'success': False,
                    'error': 'OpenAI API key not configured for image generation'
                }
            
            slides = slides_data.get('slides', [])
            generated_images = []
            
            for slide in slides:
                visual_suggestion = slide.get('visual_suggestion', '')
                slide_num = slide.get('number', 0)
                slide_title = slide.get('title', 'Slide')
                
                if not visual_suggestion or visual_suggestion == 'No visual':
                    generated_images.append({
                        'slide_num': slide_num,
                        'success': False,
                        'error': 'No visual suggestion'
                    })
                    continue
                
                # Create enhanced prompt for DALL-E with natural, artistic style
                # Include slide context for relevance
                slide_content = ' '.join(slide.get('content', []))[:100]  # First 100 chars of content
                
                dalle_prompt = (
                    f"{visual_suggestion}. "
                    f"Context: {slide_title}. "
                    f"Style: Natural photorealistic photography with artistic composition, "
                    f"professional grade image quality, natural lighting and depth of field, "
                    f"captured with professional camera aesthetics, organic and authentic feel, "
                    f"cinematic color grading, expertly composed by a professional photographer, "
                    f"high-end editorial style, real-world setting, lifelike textures and details, "
                    f"suitable for professional business presentation."
                )
                
                # Generate image using DALL-E
                image_result = self.ai_service.chat(
                    provider='dalle',
                    message=dalle_prompt,
                    api_key=api_key,
                    files=None,
                    conversation_history=[],
                    version='dall-e-3'
                )
                
                if 'error' in image_result:
                    generated_images.append({
                        'slide_num': slide_num,
                        'success': False,
                        'error': image_result.get('error')
                    })
                else:
                    # Save image data
                    generated_images.append({
                        'slide_num': slide_num,
                        'success': True,
                        'image_url': image_result.get('image_url', ''),
                        'image_path': image_result.get('image_path', ''),
                        'filename': image_result.get('filename', '')
                    })
            
            # Store images in session
            if session_id in self.sessions:
                self.sessions[session_id]['generated_images'] = generated_images
            
            return {
                'success': True,
                'images': generated_images,
                'total': len(slides),
                'generated': sum(1 for img in generated_images if img.get('success'))
            }
            
        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }
    
    def convert_to_format(self, session_id: str, format_type: str, generate_images: bool = True) -> Dict[str, Any]:
        """
        Convert the generated slides to specified format
        format_type: 'pptx', 'html', or 'json'
        generate_images: Whether to generate DALL-E images for slides
        """
        if session_id not in self.sessions:
            return {
                'status': 'error',
                'error': 'Session not found',
                'message': 'Please generate slides first',
                'debug': f'Available sessions: {list(self.sessions.keys())}'
            }
        
        slides_data = self.sessions[session_id]['slides_data']
        
        # Normalize format type
        format_lower = format_type.lower().strip()
        
        try:
            # Generate images if requested and not already generated
            if generate_images and 'generated_images' not in self.sessions[session_id]:
                image_result = self._generate_images_for_slides(slides_data, session_id)
                if not image_result.get('success'):
                    # Continue without images if generation fails
                    self.sessions[session_id]['generated_images'] = []
            
            if format_lower in ['pptx', 'powerpoint', '.pptx', 'ppt']:
                return self._create_pptx(slides_data, session_id)
            elif format_lower in ['html', 'web', 'webpage']:
                return self._create_html(slides_data, session_id)
            elif format_lower in ['json', 'data']:
                return {
                    'status': 'success',
                    'format': 'json',
                    'data': slides_data,
                    'message': 'JSON data ready for download'
                }
            else:
                return {
                    'status': 'error',
                    'error': f'Invalid format: "{format_type}"',
                    'message': f'Supported formats: PowerPoint (pptx), HTML (html), JSON (json). Received: "{format_type}"',
                    'debug': f'Normalized format: "{format_lower}"'
                }
                
        except Exception as e:
            return {
                'status': 'error',
                'error': str(e),
                'message': f'Failed to convert: {str(e)}'
            }
    
    def _create_pptx(self, slides_data: Dict[str, Any], session_id: str) -> Dict[str, Any]:
        """Create PowerPoint file with embedded images"""
        if not PPTX_AVAILABLE:
            return {
                'status': 'error',
                'error': 'python-pptx not available',
                'message': 'PowerPoint generation requires python-pptx library'
            }
        
        try:
            prs = Presentation()
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)
            
            # Get generated images if available
            generated_images = self.sessions.get(session_id, {}).get('generated_images', [])
            images_by_slide = {img['slide_num']: img for img in generated_images if img.get('success')}
            
            # Title slide
            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            title.text = slides_data.get('title', 'Presentation')
            subtitle.text = f"Generated by PresentMate with AI-Generated Visuals"
            
            # Content slides
            for slide_data in slides_data.get('slides', []):
                # Use blank layout if we have an image, otherwise use bullet layout
                slide_num = slide_data.get('number', 0)
                has_image = slide_num in images_by_slide
                
                if has_image:
                    # Blank layout for custom positioning
                    blank_layout = prs.slide_layouts[6]
                    slide = prs.slides.add_slide(blank_layout)
                    
                    # Add title
                    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
                    title_frame = title_box.text_frame
                    title_frame.text = slide_data.get('title', 'Slide')
                    title_frame.paragraphs[0].font.size = Pt(32)
                    title_frame.paragraphs[0].font.bold = True
                    
                    # Add image
                    image_info = images_by_slide[slide_num]
                    image_path = image_info.get('image_path')
                    if image_path and os.path.exists(image_path):
                        # Image on right side
                        slide.shapes.add_picture(image_path, Inches(5.2), Inches(1.5), width=Inches(4.3))
                    
                    # Add content bullets on left side
                    text_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(5.5))
                    text_frame = text_box.text_frame
                    text_frame.word_wrap = True
                    
                    for point in slide_data.get('content', []):
                        p = text_frame.add_paragraph()
                        p.text = point
                        p.level = 0
                        p.font.size = Pt(16)
                        p.space_before = Pt(12)
                else:
                    # Standard bullet layout
                    bullet_slide_layout = prs.slide_layouts[1]
                    slide = prs.slides.add_slide(bullet_slide_layout)
                    shapes = slide.shapes
                    
                    title_shape = shapes.title
                    body_shape = shapes.placeholders[1]
                    
                    title_shape.text = slide_data.get('title', 'Slide')
                    
                    tf = body_shape.text_frame
                    for point in slide_data.get('content', []):
                        p = tf.add_paragraph()
                        p.text = point
                        p.level = 0
            
            # Save file
            output_dir = 'generated_presentations'
            os.makedirs(output_dir, exist_ok=True)
            filename = f"presentation_{session_id}_{int(datetime.now().timestamp())}.pptx"
            filepath = os.path.join(output_dir, filename)
            prs.save(filepath)
            
            images_count = len(images_by_slide)
            message = f'PowerPoint file created successfully with {images_count} AI-generated images!' if images_count > 0 else 'PowerPoint file created successfully!'
            
            return {
                'status': 'success',
                'format': 'pptx',
                'file_path': filepath,
                'filename': filename,
                'message': message,
                'download_url': f'/api/presentmate/download/{filename}',
                'images_embedded': images_count
            }
            
        except Exception as e:
            return {
                'status': 'error',
                'error': str(e),
                'message': f'Failed to create PowerPoint: {str(e)}'
            }
    
    def _create_html(self, slides_data: Dict[str, Any], session_id: str) -> Dict[str, Any]:
        """Create HTML presentation with embedded images and save as downloadable file"""
        title = slides_data.get('title', 'Presentation')
        slides = slides_data.get('slides', [])
        
        # Get generated images if available
        generated_images = self.sessions.get(session_id, {}).get('generated_images', [])
        images_by_slide = {img['slide_num']: img for img in generated_images if img.get('success')}
        
        html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <style>
        body {{
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            margin: 0;
            padding: 20px;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        }}
        .container {{
            max-width: 1200px;
            margin: 0 auto;
        }}
        .slide {{
            background: white;
            border-radius: 12px;
            padding: 40px;
            margin: 20px 0;
            box-shadow: 0 10px 30px rgba(0,0,0,0.2);
        }}
        .slide-title {{
            color: #667eea;
            font-size: 32px;
            font-weight: bold;
            margin-bottom: 20px;
            border-bottom: 3px solid #667eea;
            padding-bottom: 10px;
        }}
        .slide-content {{
            font-size: 18px;
            line-height: 1.8;
            display: flex;
            gap: 30px;
            align-items: flex-start;
        }}
        .slide-text {{
            flex: 1;
        }}
        .slide-content li {{
            margin: 10px 0;
        }}
        .slide-image {{
            flex: 0 0 400px;
            max-width: 400px;
        }}
        .slide-image img {{
            width: 100%;
            height: auto;
            border-radius: 8px;
            box-shadow: 0 4px 15px rgba(0,0,0,0.2);
        }}
        .title-slide {{
            text-align: center;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
        }}
        .title-slide h1 {{
            color: white;
            font-size: 48px;
        }}
        .visual-note {{
            color: #888;
            font-style: italic;
            font-size: 14px;
            margin-top: 15px;
        }}
        @media (max-width: 768px) {{
            .slide-content {{
                flex-direction: column;
            }}
            .slide-image {{
                flex: 1;
                max-width: 100%;
            }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="slide title-slide">
            <h1>{title}</h1>
            <p>Generated by PresentMate with AI-Generated Visuals</p>
        </div>
"""
        
        for slide in slides:
            slide_num = slide.get('number', 0)
            slide_title = slide.get('title', 'Slide')
            content = slide.get('content', [])
            visual = slide.get('visual_suggestion', '')
            has_image = slide_num in images_by_slide
            
            html += f"""
        <div class="slide">
            <div class="slide-title">Slide {slide_num}: {slide_title}</div>
            <div class="slide-content">
                <div class="slide-text">
                    <ul>
"""
            for point in content:
                html += f"                        <li>{point}</li>\n"
            
            html += """                    </ul>
"""
            if visual and not has_image:
                html += f'                    <div class="visual-note">🎨 Suggested Visual: {visual}</div>\n'
            
            html += """                </div>
"""
            
            # Add image if available
            if has_image:
                image_info = images_by_slide[slide_num]
                image_url = image_info.get('image_url', '')
                html += f"""                <div class="slide-image">
                    <img src="{image_url}" alt="{visual}">
                    <div class="visual-note" style="text-align: center; margin-top: 10px;">AI-Generated Visual</div>
                </div>
"""
            
            html += """            </div>
        </div>
"""
        
        html += """    </div>
</body>
</html>"""
        
        # Save HTML file to disk for download
        try:
            output_dir = 'generated_presentations'
            os.makedirs(output_dir, exist_ok=True)
            filename = f"presentation_{int(datetime.now().timestamp())}.html"
            filepath = os.path.join(output_dir, filename)
            
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(html)
            
            images_count = len(images_by_slide)
            message = f'HTML presentation created with {images_count} AI-generated images!' if images_count > 0 else 'HTML presentation created!'
            
            return {
                'status': 'success',
                'format': 'html',
                'html': html,
                'file_path': filepath,
                'filename': filename,
                'message': message,
                'download_url': f'/api/presentmate/download/{filename}',
                'images_embedded': images_count
            }
        except Exception as e:
            # Fallback: return HTML content even if file save fails
            return {
                'status': 'success',
                'format': 'html',
                'html': html,
                'message': 'HTML presentation created!',
                'note': f'File save failed: {str(e)}'
            }
