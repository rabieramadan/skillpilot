"""
AI-powered Question Generator Service
Generates bilingual (Arabic/English) questions from uploaded files
Supports True/False and Multiple Choice questions
"""

import json
import re
from typing import Dict, List, Any, Optional
from datetime import datetime

from app.services.ai_transport import extract_json

#: How much source text to send. Well under the smallest context window in the
#: catalogue, while still covering a full lecture or chapter.
MAX_SOURCE_CHARS = 60_000


class QuestionGenerator:
    """Generate exam and survey questions from uploaded content using AI"""
    
    def __init__(self, ai_service, config):
        self.ai_service = ai_service
        self.config = config
    
    def extract_text_from_file(self, file_content: bytes, file_type: str) -> str:
        """Extract text content from uploaded file"""
        text = ""
        
        try:
            if file_type in ['txt', 'text']:
                text = file_content.decode('utf-8', errors='ignore')
            
            elif file_type == 'pdf':
                try:
                    import pdfplumber
                    import io
                    with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                        for page in pdf.pages:
                            page_text = page.extract_text()
                            if page_text:
                                text += page_text + "\n"
                except Exception as e:
                    print(f"PDF extraction error: {e}")
            
            elif file_type in ['doc', 'docx']:
                try:
                    from docx import Document
                    import io
                    doc = Document(io.BytesIO(file_content))
                    for para in doc.paragraphs:
                        text += para.text + "\n"
                except Exception as e:
                    print(f"DOCX extraction error: {e}")
            
            elif file_type in ['ppt', 'pptx']:
                try:
                    from pptx import Presentation
                    import io
                    prs = Presentation(io.BytesIO(file_content))
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                text += shape.text + "\n"
                except Exception as e:
                    print(f"PPTX extraction error: {e}")
            
        except Exception as e:
            print(f"File extraction error: {e}")
        
        return text.strip()
    
    def generate_questions(
        self,
        content: str,
        num_true_false: int = 5,
        num_mcq: int = 5,
        question_type: str = 'exam',
        difficulty: str = 'medium',
        topic: str = None
    ) -> Dict[str, Any]:
        """
        Generate bilingual questions from content using AI
        
        Args:
            content: Text content to generate questions from
            num_true_false: Number of True/False questions
            num_mcq: Number of Multiple Choice questions
            question_type: 'exam' or 'survey'
            difficulty: 'easy', 'medium', or 'hard'
            topic: Optional topic/subject area
        
        Returns:
            Dictionary with 'questions' list and 'success' status
        """
        
        if not content or len(content.strip()) < 50:
            return {'success': False, 'error': 'Content too short for question generation'}
        
        # Current models have context windows measured in hundreds of
        # thousands of tokens. The old 8,000-character cut discarded most of a
        # lecture PDF, so questions were generated from the first few pages.
        content_preview = content[:MAX_SOURCE_CHARS]
        
        prompt = self._build_generation_prompt(
            content_preview,
            num_true_false,
            num_mcq,
            question_type,
            difficulty,
            topic
        )
        
        providers_to_try = []
        
        openai_key = getattr(self.config, 'OPENAI_API_KEY', None)
        grok_key = getattr(self.config, 'GROK_API_KEY', None)
        deepseek_key = getattr(self.config, 'DEEPSEEK_API_KEY', None)
        claude_key = getattr(self.config, 'ANTHROPIC_API_KEY', None)
        
        if openai_key:
            providers_to_try.append(('openai', openai_key, None))
        if grok_key:
            providers_to_try.append(('grok', grok_key, None))
        if deepseek_key:
            providers_to_try.append(('deepseek', deepseek_key, None))
        if claude_key:
            providers_to_try.append(('claude', claude_key, None))
        
        if not providers_to_try:
            return {'success': False, 'error': 'No AI API key configured'}
        
        last_error = None
        
        for provider, api_key, version in providers_to_try:
            try:
                print(f"Trying question generation with {provider}...")
                
                response = self.ai_service.chat(
                    provider=provider,
                    message=prompt,
                    api_key=api_key,
                    version=version
                )
                
                if 'error' in response:
                    last_error = response['error']
                    print(f"{provider} failed: {last_error}")
                    continue
                
                questions = self._parse_ai_response(response.get('text', ''))
                
                if not questions:
                    last_error = 'Failed to parse generated questions'
                    print(f"{provider} returned unparseable response")
                    continue
                
                print(f"Successfully generated {len(questions)} questions with {provider}")
                return {'success': True, 'questions': questions, 'provider': provider}
                
            except Exception as e:
                last_error = str(e)
                print(f"{provider} exception: {e}")
                continue
        
        return {'success': False, 'error': f'All providers failed. Last error: {last_error}'}
    
    def _build_generation_prompt(
        self,
        content: str,
        num_true_false: int,
        num_mcq: int,
        question_type: str,
        difficulty: str,
        topic: str
    ) -> str:
        """Build the AI prompt for question generation"""
        
        topic_str = f" about {topic}" if topic else ""
        
        prompt = f"""You write assessment items for a university course. Write
exam questions{topic_str} from the source material below.

WHAT TO PRODUCE:
1. Exactly {num_true_false} true/false questions
2. Exactly {num_mcq} multiple-choice questions, four options each (A, B, C, D)
3. Difficulty: {difficulty}
4. Every question in both English and Arabic
5. A correct answer for every question

RULES THAT DECIDE WHETHER THESE QUESTIONS ARE USABLE:
- Every question must be answerable from the source material alone, and the
  correct answer must be verifiable by pointing at a specific passage in it.
  Do not draw on outside knowledge and do not invent facts, figures or names
  that the material does not contain.
- Test understanding, not recall of a sentence's wording. A student who
  understood the material should answer correctly; one who only skimmed it
  should not.
- Distractors must be plausible to someone who half-learned the material —
  a common misconception, a neighbouring concept, a plausible wrong number.
  Never use "all of the above", "none of the above", joke options, or options
  that are obviously the wrong length or grammatical shape.
- Spread the correct answers across A, B, C and D roughly evenly. Do not put
  most of them in one position.
- Roughly half the true/false statements should be false, and the false ones
  must be plausible.
- Each question must stand alone. Do not write "as mentioned above", do not
  refer to figure or page numbers, and do not repeat a question you have
  already asked in different words.
- The Arabic must be a natural translation that a native speaker would write,
  not a word-for-word transliteration. Keep technical terms, code and
  mathematical notation in their original form.
- Keep each question under 40 words and each option under 15.

OUTPUT FORMAT (JSON):
Return a valid JSON array with this exact structure:
[
  {{
    "type": "true_false",
    "question_en": "English question text",
    "question_ar": "Arabic question text (نص السؤال بالعربية)",
    "correct_answer": "true" or "false",
    "difficulty": "{difficulty}"
  }},
  {{
    "type": "multiple_choice",
    "question_en": "English question text",
    "question_ar": "Arabic question text (نص السؤال بالعربية)",
    "options_en": ["Option A", "Option B", "Option C", "Option D"],
    "options_ar": ["الخيار أ", "الخيار ب", "الخيار ج", "الخيار د"],
    "correct_answer": "A", "B", "C", or "D",
    "difficulty": "{difficulty}"
  }}
]

CONTENT TO GENERATE QUESTIONS FROM:
---
{content}
---

Return exactly {num_true_false} true/false and {num_mcq} multiple-choice
questions as a JSON array, and nothing else — no prose before or after it, no
code fence, no explanation."""
        
        return prompt
    
    def _parse_ai_response(self, response_text: str) -> List[Dict[str, Any]]:
        """Parse the AI response and extract questions"""
        
        questions = extract_json(response_text, expect='array')
        if questions is None:
            print('Question generator: the model did not return a JSON array.')
            return []

        validated_questions = []
        for question in questions:
            if not isinstance(question, dict):
                continue
            question = self._sanitize_question(question)
            if self._validate_question(question):
                validated_questions.append(question)
        return validated_questions
    
    def _sanitize_question(self, question: Dict[str, Any]) -> Dict[str, Any]:
        """Sanitize question data - fix common AI formatting issues.

        Handles every shape we have observed AI providers return for the
        options field: a JSON list of strings (good), a single newline-
        joined string, a list containing one big newline-joined string,
        a list of dicts like {"A": "..."} or {"label": "A", "text": "..."},
        and stray "A) " / "1. " / leading-bullet prefixes.
        """
        import re

        def _flatten_one(opt):
            # dict like {"A": "text"} or {"label":"A","text":"..."} -> "text"
            if isinstance(opt, dict):
                if 'text' in opt:
                    return [str(opt['text']).strip()]
                if 'value' in opt:
                    return [str(opt['value']).strip()]
                if len(opt) == 1:
                    return [str(list(opt.values())[0]).strip()]
                return [str(opt)]
            if not isinstance(opt, str):
                return [str(opt).strip()]
            # split on real or escaped newlines
            parts = re.split(r'\r?\n|\\n', opt)
            cleaned = []
            for p in parts:
                p = p.strip()
                if not p:
                    continue
                # strip leading "A) ", "A. ", "A: ", "1) ", "- ", "* "
                p = re.sub(r'^\s*(?:[A-Ea-e]|[0-9]+)\s*[\)\.\:\-]\s+', '', p)
                p = re.sub(r'^\s*[\-\*\u2022]\s+', '', p)
                if p:
                    cleaned.append(p)
            return cleaned

        for key in ['options_en', 'options_ar']:
            if key not in question:
                continue
            opts = question[key]
            if opts is None:
                question[key] = []
                continue
            if isinstance(opts, str):
                question[key] = _flatten_one(opts)
                continue
            if isinstance(opts, list):
                flat = []
                for opt in opts:
                    flat.extend(_flatten_one(opt))
                question[key] = flat
        
        for key in ['question_en', 'question_ar']:
            if key in question and isinstance(question[key], str):
                question[key] = question[key].strip()
        
        return question
    
    def _validate_question(self, question: Dict[str, Any]) -> bool:
        """Validate a parsed question has required fields"""
        
        required_fields = ['type', 'question_en', 'correct_answer']
        
        for field in required_fields:
            if field not in question:
                return False
        
        if question['type'] == 'multiple_choice':
            if 'options_en' not in question or len(question.get('options_en', [])) < 2:
                return False
        
        return True
    
    def generate_feedback_survey_questions(self, language: str = 'both') -> List[Dict[str, Any]]:
        """
        Generate generic feedback survey questions for evaluating:
        - Course materials
        - Teacher/Instructor
        - Teaching venue/place
        
        Returns pre-defined bilingual feedback questions
        """
        
        feedback_questions = [
            {
                "category": "course_materials",
                "question_en": "How would you rate the quality of the course materials?",
                "question_ar": "كيف تقيم جودة مواد الدورة؟",
                "type": "rating",
                "options_en": ["Poor", "Fair", "Good", "Very Good", "Excellent"],
                "options_ar": ["ضعيف", "مقبول", "جيد", "جيد جداً", "ممتاز"]
            },
            {
                "category": "course_materials",
                "question_en": "The course content was well-organized and easy to follow.",
                "question_ar": "كان محتوى الدورة منظماً جيداً وسهل المتابعة.",
                "type": "likert",
                "options_en": ["Strongly Disagree", "Disagree", "Neutral", "Agree", "Strongly Agree"],
                "options_ar": ["أرفض بشدة", "أرفض", "محايد", "أوافق", "أوافق بشدة"]
            },
            {
                "category": "course_materials",
                "question_en": "The learning resources (videos, documents, exercises) were helpful.",
                "question_ar": "كانت مصادر التعلم (الفيديوهات، المستندات، التمارين) مفيدة.",
                "type": "likert",
                "options_en": ["Strongly Disagree", "Disagree", "Neutral", "Agree", "Strongly Agree"],
                "options_ar": ["أرفض بشدة", "أرفض", "محايد", "أوافق", "أوافق بشدة"]
            },
            {
                "category": "teacher",
                "question_en": "How would you rate the instructor's knowledge of the subject?",
                "question_ar": "كيف تقيم معرفة المدرب بالموضوع؟",
                "type": "rating",
                "options_en": ["Poor", "Fair", "Good", "Very Good", "Excellent"],
                "options_ar": ["ضعيف", "مقبول", "جيد", "جيد جداً", "ممتاز"]
            },
            {
                "category": "teacher",
                "question_en": "The instructor explained concepts clearly and effectively.",
                "question_ar": "شرح المدرب المفاهيم بوضوح وفعالية.",
                "type": "likert",
                "options_en": ["Strongly Disagree", "Disagree", "Neutral", "Agree", "Strongly Agree"],
                "options_ar": ["أرفض بشدة", "أرفض", "محايد", "أوافق", "أوافق بشدة"]
            },
            {
                "category": "teacher",
                "question_en": "The instructor was responsive to questions and provided helpful feedback.",
                "question_ar": "كان المدرب متجاوباً مع الأسئلة وقدم ملاحظات مفيدة.",
                "type": "likert",
                "options_en": ["Strongly Disagree", "Disagree", "Neutral", "Agree", "Strongly Agree"],
                "options_ar": ["أرفض بشدة", "أرفض", "محايد", "أوافق", "أوافق بشدة"]
            },
            {
                "category": "teacher",
                "question_en": "The instructor created an engaging learning environment.",
                "question_ar": "أنشأ المدرب بيئة تعليمية جذابة.",
                "type": "likert",
                "options_en": ["Strongly Disagree", "Disagree", "Neutral", "Agree", "Strongly Agree"],
                "options_ar": ["أرفض بشدة", "أرفض", "محايد", "أوافق", "أوافق بشدة"]
            },
            {
                "category": "venue",
                "question_en": "How would you rate the training venue/facility?",
                "question_ar": "كيف تقيم مكان/منشأة التدريب؟",
                "type": "rating",
                "options_en": ["Poor", "Fair", "Good", "Very Good", "Excellent"],
                "options_ar": ["ضعيف", "مقبول", "جيد", "جيد جداً", "ممتاز"]
            },
            {
                "category": "venue",
                "question_en": "The training room was comfortable and conducive to learning.",
                "question_ar": "كانت قاعة التدريب مريحة ومناسبة للتعلم.",
                "type": "likert",
                "options_en": ["Strongly Disagree", "Disagree", "Neutral", "Agree", "Strongly Agree"],
                "options_ar": ["أرفض بشدة", "أرفض", "محايد", "أوافق", "أوافق بشدة"]
            },
            {
                "category": "venue",
                "question_en": "The technical equipment (computers, projectors, internet) worked well.",
                "question_ar": "عملت المعدات التقنية (الحواسيب، أجهزة العرض، الإنترنت) بشكل جيد.",
                "type": "likert",
                "options_en": ["Strongly Disagree", "Disagree", "Neutral", "Agree", "Strongly Agree"],
                "options_ar": ["أرفض بشدة", "أرفض", "محايد", "أوافق", "أوافق بشدة"]
            },
            {
                "category": "overall",
                "question_en": "Overall, how satisfied are you with this training course?",
                "question_ar": "بشكل عام، ما مدى رضاك عن هذه الدورة التدريبية؟",
                "type": "rating",
                "options_en": ["Very Dissatisfied", "Dissatisfied", "Neutral", "Satisfied", "Very Satisfied"],
                "options_ar": ["غير راضٍ جداً", "غير راضٍ", "محايد", "راضٍ", "راضٍ جداً"]
            },
            {
                "category": "overall",
                "question_en": "Would you recommend this course to others?",
                "question_ar": "هل توصي بهذه الدورة للآخرين؟",
                "type": "likert",
                "options_en": ["Definitely Not", "Probably Not", "Maybe", "Probably Yes", "Definitely Yes"],
                "options_ar": ["بالتأكيد لا", "على الأرجح لا", "ربما", "على الأرجح نعم", "بالتأكيد نعم"]
            },
            {
                "category": "overall",
                "question_en": "What did you like most about the course?",
                "question_ar": "ما أكثر ما أعجبك في الدورة؟",
                "type": "text",
                "options_en": [],
                "options_ar": []
            },
            {
                "category": "overall",
                "question_en": "What suggestions do you have for improving this course?",
                "question_ar": "ما هي اقتراحاتك لتحسين هذه الدورة؟",
                "type": "text",
                "options_en": [],
                "options_ar": []
            }
        ]
        
        return feedback_questions


    # ----------------------------------------------------------------
    # Phase 3 — Spaced retention quiz generator
    # ----------------------------------------------------------------
    def generate_retention_quiz(self, skill_name: str, skill_name_ar: str = None,
                                num_questions: int = 5, language: str = "en") -> Dict[str, Any]:
        """Generate a short bilingual retention quiz for a previously-mastered skill.

        Falls back to a deterministic template-based quiz when AI is unavailable
        so retention scheduling never breaks silently.
        """
        prompt = (
            f"Generate {num_questions} short retention-check questions for the skill "
            f"\"{skill_name}\". Mix true/false and multiple choice. Return JSON: "
            "{\"questions\":[{\"type\":\"multiple_choice|true_false\","
            "\"question\":\"...\",\"question_ar\":\"...\","
            "\"options\":[..],\"options_ar\":[..],"
            "\"correct_answer\":\"...\",\"correct_answer_ar\":\"...\"}]}"
        )
        try:
            if self.ai_service and hasattr(self.ai_service, "chat"):
                resp = self.ai_service.chat("openai", prompt, "", [], [], "", language)
                text = (resp or {}).get("text") or ""
                match = re.search(r"\{[\s\S]*\}", text)
                if match:
                    data = json.loads(match.group(0))
                    qs = data.get("questions") or []
                    if qs:
                        return {"success": True, "source": "ai",
                                "questions": qs[:num_questions]}
        except Exception as e:
            print(f"AI retention quiz generation failed: {e}")

        # Deterministic fallback
        ar = skill_name_ar or skill_name
        questions = []
        for i in range(num_questions):
            questions.append({
                "type": "true_false",
                "question": f"Retention check {i+1}: \"{skill_name}\" still applies in this scenario.",
                "question_ar": f"اختبار استرجاع {i+1}: \"{ar}\" لا يزال ينطبق في هذا السياق.",
                "options": ["True", "False"],
                "options_ar": ["صحيح", "خطأ"],
                "correct_answer": "True",
                "correct_answer_ar": "صحيح",
                "topic": skill_name,
                "topic_ar": ar,
                "difficulty": "easy",
            })
        return {"success": True, "source": "fallback", "questions": questions}
